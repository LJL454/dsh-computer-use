"""Acceptance for DSH故障定位与修复方案.md §七.2 items 2-4.

The two presets now point `dsh-tool-workspace-dependencies` at the installed runtime root, so the
thing that has to be true is not "the config validates" but "the interpreters it hands back really
run". This reads a small CSV with pandas, then writes and re-reads a DOCX and an XLSX with
python-docx / openpyxl — the three libraries the preset's description promises.

It uses only the interpreter path the tool itself reports, so a wrong `source` shows up here as a
failure rather than as a surprise in a user's task.
"""
import csv
import json
import os
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from docx import Document

result = {"python": sys.executable, "version": sys.version.split()[0], "checks": []}


def check(name, ok, detail=""):
    result["checks"].append({"name": name, "ok": bool(ok), "detail": str(detail)})


with tempfile.TemporaryDirectory() as tmp:
    # 3. a small CSV through pandas
    csv_path = os.path.join(tmp, "sales.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["region", "units", "price"])
        writer.writerows([
            ["north", 3, 10.0],
            ["south", 5, 12.5],
            ["north", 2, 10.0],
            ["east", 7, 9.0],
        ])

    import pandas as pd

    frame = pd.read_csv(csv_path)
    frame["revenue"] = frame["units"] * frame["price"]
    grouped = frame.groupby("region", as_index=False)["revenue"].sum().sort_values("region")
    result["pandas"] = {
        "version": pd.__version__,
        "rows": int(len(frame)),
        "total": float(frame["revenue"].sum()),
        "by_region": {row.region: float(row.revenue) for row in grouped.itertuples()},
    }
    check("pandas read the CSV", len(frame) == 4, f"rows={len(frame)}")
    # north 3x10.0 + 2x10.0 = 50.0, south 5x12.5 = 62.5, east 7x9.0 = 63.0
    check("pandas computed the summary", abs(frame["revenue"].sum() - 175.5) < 1e-9, str(frame["revenue"].sum()))
    check("pandas grouped by region", list(grouped["region"]) == ["east", "north", "south"], list(grouped["region"]))
    check(
        "pandas grouped the right amounts",
        [float(value) for value in grouped["revenue"]] == [63.0, 50.0, 62.5],
        list(grouped["revenue"]),
    )

    # 4a. a minimal DOCX, written and read back
    docx_path = os.path.join(tmp, "report.docx")
    document = Document()
    document.add_heading("Sales", level=1)
    document.add_paragraph(f"Total revenue: {frame['revenue'].sum():.2f}")
    for row in grouped.itertuples():
        document.add_paragraph(f"{row.region}: {row.revenue:.2f}")
    document.save(docx_path)

    reopened = Document(docx_path)
    paragraphs = [p.text for p in reopened.paragraphs if p.text.strip()]
    check("python-docx wrote a DOCX", os.path.getsize(docx_path) > 0, f"{os.path.getsize(docx_path)} bytes")
    check("python-docx read it back", paragraphs[0] == "Sales" and any("Total revenue: 175.50" in p for p in paragraphs), paragraphs[:2])

    # 4b. a minimal XLSX, written and read back
    xlsx_path = os.path.join(tmp, "report.xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "sales"
    sheet.append(["region", "revenue"])
    for row in grouped.itertuples():
        sheet.append([row.region, float(row.revenue)])
    book.save(xlsx_path)

    loaded = load_workbook(xlsx_path)
    cells = list(loaded["sales"].iter_rows(values_only=True))
    check("openpyxl wrote an XLSX", os.path.getsize(xlsx_path) > 0, f"{os.path.getsize(xlsx_path)} bytes")
    check("openpyxl read it back", cells[0] == ("region", "revenue") and len(cells) == 4, cells)
    check("the values survived the round trip", abs(cells[1][1] - 63.0) < 1e-9 and abs(cells[3][1] - 62.5) < 1e-9, cells[1:])

print(json.dumps(result, ensure_ascii=False, indent=2))
sys.exit(0 if all(item["ok"] for item in result["checks"]) else 1)
